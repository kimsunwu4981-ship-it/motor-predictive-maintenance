# -*- coding: utf-8 -*-
"""sim_vs_analysis.py — Isaac Sim 검증 로그(logs/verify/*) ↔ 해석 기준(data/electrical/analysis_reference.json) 대조 (V2) + 그림 (V4). 수기 숫자 없음.

읽는 것: logs/verify/electrical_a{0,90,180}/{state,events}.csv·meta.json·waveform_*.csv, logs/verify/thermal_itsc_{4,12,16}turn/, analysis_reference.json, recompute_i2i1.py(독립 재계산)
쓰는 것: docs/verification/sim_vs_analysis.json(모든 행), sim_vs_analysis.md(표), sim_vs_analysis.xlsx(수식 판정), docs/verification/img/*.png, logs/verify/export_<시각>.csv(V2 형식)

허용(지시서 V2 표): I₂/I₁ 결함분 ±1 % · 합성값 범위 안 · I₂/I_n ±1 % · 불평형 ±2 % · I_f MCCM ±1 % · 손실 ±2 % · 권선 ΔT ±5 % · 판정 일치 · 정상 오경보 0 · 독립 재계산 ±0.5 %
  py -3.12 -X utf8 tools/verify/sim_vs_analysis.py [--logs logs/verify] [--no-fig]
"""
from __future__ import annotations

import argparse
import cmath
import csv
import glob
import json
import math
import os
import statistics as st
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
VIZ = os.path.abspath(os.path.join(HERE, "..", ".."))
REPO = os.path.abspath(os.path.join(VIZ, ".."))
OUT_DIR = os.path.join(REPO, "docs", "verification")
IMG_DIR = os.path.join(OUT_DIR, "img")
sys.path.insert(0, HERE)
from recompute_i2i1 import analyze as recompute  # noqa: E402  (엔진 비의존)

TOL = {"i2i1_fault": 0.01, "i2i1_composite": None, "i2_over_in": 0.01, "unbalance": 0.02, "i_f": 0.01, "p_loop": 0.02, "dT": 0.05, "recompute": 0.005}
ANGLES = (0, 90, 180)
STEADY_S = 5.0


def _py(o):
    """numpy 스칼라(np.bool_/float64) → 파이썬 기본형 (json 직렬화)."""
    if hasattr(o, "item"):
        return o.item()
    raise TypeError(str(type(o)))


def read_csv(p):
    return list(csv.DictReader(open(p, encoding="utf-8", newline="")))


def fnum(x, default=None):
    try:
        return float(x)
    except (TypeError, ValueError):
        return default


def load_run(d):
    if not os.path.exists(os.path.join(d, "state.csv")):
        return None
    return {"dir": d, "state": read_csv(os.path.join(d, "state.csv")), "events": read_csv(os.path.join(d, "events.csv")),
            "meta": json.load(open(os.path.join(d, "meta.json"), encoding="utf-8")), "waveforms": sorted(glob.glob(os.path.join(d, "waveform_*.csv")))}


def steady(rows, step, t0, t1, key):
    seg = [fnum(r[key]) for r in rows if r["step"] == step and t1 - STEADY_S <= fnum(r["t_real_s"]) < t1 and fnum(r[key]) is not None]
    if not seg:
        return None
    m = st.mean(seg); s = st.pstdev(seg) if len(seg) > 1 else 0.0
    return {"mean": m, "std": s, "std_pct": (s / m * 100) if m else 0.0, "n": len(seg), "steady": (s / m * 100 < 1.0) if m else True}


def pct(a, b):
    return None if (a is None or b in (None, 0)) else (a / b - 1) * 100


def row(item, scen, sim, ref, tol, unit="", note="", src="", kind="rel", extra=None):
    """kind: rel(±tol 상대) | range([lo,hi]) | match(문자열 일치) | bool"""
    if kind == "rel":
        if isinstance(ref, (int, float)) and ref == 0 and isinstance(sim, (int, float)):      # 기준 0(정상 결함분 등): 절대값 1e-6 안이면 일치
            ok = abs(sim) <= 1e-6; dev = None
        else:
            d = pct(sim, ref); ok = d is not None and abs(d) <= tol * 100
            dev = None if d is None else round(d, 3)
    elif kind == "range":
        ok = sim is not None and ref[0] - 1e-9 <= sim <= ref[1] + 1e-9; dev = None
    elif kind == "match":
        ok = sim == ref; dev = None
    else:
        ok = bool(sim); dev = None
    return {"item": item, "scenario": scen, "sim": sim, "ref": ref, "tolerance": tol, "dev_pct": dev, "pass": bool(ok), "unit": unit, "note": note, "source": src, "kind": kind, **(extra or {})}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--logs", default=os.path.join(VIZ, "logs", "verify"))
    ap.add_argument("--no-fig", action="store_true")
    a = ap.parse_args()
    ref = json.load(open(os.path.join(VIZ, "data", "electrical", "analysis_reference.json"), encoding="utf-8"))
    S = ref["scenarios"]; h = float(ref["inherent_i2i1"]); I_n = float(ref["I_n_A"]); warn, danger = ref["thresholds"]["i2i1_warn"], ref["thresholds"]["i2i1_danger"]
    runs = {f"a{ang}": load_run(os.path.join(a.logs, f"electrical_a{ang}")) for ang in ANGLES}
    thermal = {T: load_run(os.path.join(a.logs, f"thermal_itsc_{T}turn")) for T in (4, 12, 16)}
    base = runs["a0"]
    if base is None:
        raise SystemExit("logs/verify/electrical_a0 없음 — Isaac 배치를 먼저 실행")
    rows_out, sens, events_tbl, recompute_tbl, timeseries = [], [], [], [], {}
    scene_t = {s["id"]: (float(s["t0_s"]), float(s["t0_s"]) + float(s["duration_s"]), s) for s in base["meta"]["steps"]}

    # ---------------- 전기 시나리오별 (각도 0 실행)
    for sid, (t0, t1, step) in scene_t.items():
        R = S.get(sid); case = step["electrical_case"]
        st_i2 = steady(base["state"], sid, t0, t1, "i2_over_i1"); st_ang = steady(base["state"], sid, t0, t1, "i2_angle_deg")
        st_in = steady(base["state"], sid, t0, t1, "i2_over_in"); st_if = steady(base["state"], sid, t0, t1, "i_f_A"); st_pl = steady(base["state"], sid, t0, t1, "p_loop_total_W")
        rms = [steady(base["state"], sid, t0, t1, k) for k in ("Ia_rms", "Ib_rms", "Ic_rms")]
        timeseries[sid] = {"t": [fnum(r["t_real_s"]) for r in base["state"] if r["step"] == sid], "i2i1": [fnum(r["i2_over_i1"]) for r in base["state"] if r["step"] == sid],
                           "status": [r["status"] for r in base["state"] if r["step"] == sid], "t0": t0, "t1": t1}
        if st_i2 is None:
            continue
        steady_ok = st_i2["steady"]
        rows_out.append(row("정상 상태 (I₂/I₁ 마지막 5 s 표준편차 < 1 %)", sid, round(st_i2["std_pct"], 4), 1.0, None, "%", f"n={st_i2['n']}", "state.csv", "bool", {"pass_override": steady_ok}))
        rows_out[-1]["pass"] = steady_ok; rows_out[-1]["sim"] = round(st_i2["std_pct"], 4)
        if R is None:                      # LUT 밖(legacy) 장면: 해석 기준 없음 — 판정만 기록
            continue
        # I₂/I₁ 결함분: 합성 페이저 − 고유 비대칭 페이저
        comp = st_i2["mean"]; ang = math.radians(st_ang["mean"] if st_ang else 0.0)
        h_ang = math.radians(float(base["meta"].get("inherent", {}).get("angle_deg", 0.0)))
        fault = abs(comp * cmath.exp(1j * ang) - h * cmath.exp(1j * h_ang)) if h else comp
        rows_out.append(row("I₂/I₁ 결함분 (합성 − 고유 비대칭 페이저)", sid, round(fault, 6), R["i2_over_i1_fault"]["value"], TOL["i2i1_fault"], "", f"합성 {comp:.5f} ∠{math.degrees(ang):.1f}°, 고유 {h} ∠{math.degrees(h_ang):.0f}°", "LUT v2 격자 (mccm_results)"))
        rows_out.append(row("I₂/I₁ 합성값 (범위 |r−h| ~ r+h)", sid, round(comp, 6), R["i2_over_i1_composite"]["range"], None, "", f"기대(각도 0°) {R['i2_over_i1_composite']['expected_by_angle']['angle_0']}", "LUT + inherent", "range"))
        rows_out.append(row("I₂/I₁ 합성값 vs 기대(각도 0°)", sid, round(comp, 6), R["i2_over_i1_composite"]["expected_by_angle"]["angle_0"], TOL["i2i1_fault"], "", "", "LUT + inherent"))
        st_i1 = steady(base["state"], sid, t0, t1, "i1_A")
        if st_in and st_i1:
            fault_in = fault * st_i1["mean"] / I_n               # D3: I₁ = LUT 정상분 RMS → 결함분 I₂/I_n = 결함분 × I₁/I_n (LUT 와 같은 정의)
            rows_out.append(row("I₂/I_n 결함분 (I₂_fault/I_n, 시뮬 I₁ 기준) vs LUT I₂/I_n", sid, round(fault_in, 6), R["i2_over_in"]["lut"], TOL["i2_over_in"], "",
                                f"시뮬 I₁ {st_i1['mean']:.4f} A (LUT 정상분 RMS, D3), LUT I₁ {R['i2_over_in']['lut_I1_A']} A", "LUT v2"))
            rows_out.append(row("I₁ 정상분 [A] vs LUT I₁", sid, round(st_i1["mean"], 5), R["i2_over_in"]["lut_I1_A"], TOL["i2_over_in"], "A", "D3: 결함 시 상승(격자 보간)", "LUT v2"))
            rows_out.append(row("I₂/I_n 표시값 vs 시뮬 기대(합성×부하)", sid, round(st_in["mean"], 6), R["i2_over_in"]["sim_expected"], TOL["i2_over_in"], "", "표시값 자체의 일관성", "LUT + inherent"))
        if all(rms):
            m = [x["mean"] for x in rms]; mean = sum(m) / 3; unb = max(abs(v - mean) for v in m) / mean * 100
            exp = R["unbalance_max_dev_pct"].get("sim_expected")
            rows_out.append(row("상 불평형 max dev [%] vs LUT 합성 기대", sid, round(unb, 4), exp if exp is not None else R["unbalance_max_dev_pct"]["lut"], TOL["unbalance"], "%", f"LUT 자체 {R['unbalance_max_dev_pct']['lut']} % (LUT I₁ 기준)", "LUT + inherent"))
        if st_if and R["i_f_A"]["mccm"]:
            rows_out.append(row("단락 루프 전류 I_f [A] vs MCCM", sid, round(st_if["mean"], 4), R["i_f_A"]["mccm"], TOL["i_f"], "A", (f"FEM(F1 보정) {R['i_f_A']['fem_corrected']['I_f_A']:.2f} A 참고" if R["i_f_A"].get("fem_corrected") else ""), "mccm_results.json"))
            if R["i_f_A"].get("fem_corrected"):
                rows_out.append(row("I_f [A] vs 2D FEM(F1 보정) — 참고 열", sid, round(st_if["mean"], 4), R["i_f_A"]["fem_corrected"]["I_f_A"], 0.05, "A", "참고(±5 % MCCM↔FEM 정합 범위)", "f2_corrections.json F1"))
        if st_pl and R["p_loop_total_W"]["lut"]:
            rows_out.append(row("단락 루프 총손실 [W] vs I_f²(R_f+μR1_hot)+단락턴 동손 (MCCM)", sid, round(st_pl["mean"], 3), R["p_loop_total_W"]["mccm"] or R["p_loop_total_W"]["lut"], TOL["p_loop"], "W", f"공식(루프 저항항만) {R['p_loop_total_W']['formula_value']:.1f} W", "mccm_results.json"))
        # 판정·탐지 지연
        ev = [e for e in base["events"] if e["step"] == sid]
        first = {}
        for e in ev:
            if e["to"] in ("WARN", "DANGER") and e["to"] not in first:
                first[e["to"]] = fnum(e["t_real_s"])
        exp_lvl = R["expected_judgment"]["i2i1_level"]
        final = [r for r in base["state"] if r["step"] == sid][-1]
        rows_out.append(row("판정 (장면 끝 종합 상태) vs 예상", sid, final["status"], R["expected_judgment"].get("overall_expected", exp_lvl), None, "", f"원인 {final['cause_signal']} {final['cause_value']} · i2i1 예상 {exp_lvl}", "thresholds.yaml", "match"))
        delay = {k: round(v - t0, 2) for k, v in first.items()}
        events_tbl.append({"scenario": sid, "t_inject_s": t0, "first_WARN_s": first.get("WARN"), "first_DANGER_s": first.get("DANGER"), "delay_s": delay, "final": final["status"], "cause": final["cause_signal"],
                           "expected_i2i1_level": exp_lvl, "expected_delay_s": R["expected_judgment"]["detection_delay_expected_s"]})
        if R["expected_judgment"]["detection_delay_expected_s"] is not None and delay:
            lvl = "DANGER" if "DANGER" in delay else "WARN"
            rows_out.append(row("탐지 지연 [s] (주입 → 첫 상향, 지속 3 s 규칙)", sid, delay.get(lvl), R["expected_judgment"]["detection_delay_expected_s"], 0.2, "s", f"첫 {lvl}", "thresholds.yaml persistence_s", "rel"))
    # 정상 오경보
    healthy_rows = [r for r in base["state"] if r["step"] == "healthy"]
    n_bad = sum(1 for r in healthy_rows if r["status"] != "SAFE")
    rows_out.append(row("정상 시나리오 오경보 건수 (SAFE 유지)", "healthy", n_bad, 0, None, "건", f"{len(healthy_rows)} 행", "state.csv", "match"))

    # ---------------- 독립 재계산 (파형 CSV)
    for run_key, run in runs.items():
        if run is None:
            continue
        for wf in run["waveforms"]:
            name = os.path.basename(wf); parts = name[:-4].split("_", 3)
            t_s = float(parts[1]); tag = parts[2]; case = parts[3]
            rc = recompute(wf, 60.0, I_n)
            cand = [r for r in run["state"] if r["case"] == case]
            if not cand:
                continue
            eng = min(cand, key=lambda r: abs(fnum(r["t_session_s"]) - t_s))
            d = pct(rc["i2_over_i1"], fnum(eng["i2_over_i1"]))
            recompute_tbl.append({"run": run_key, "file": name, "case": case, "tag": tag, "recomputed": rc["i2_over_i1"], "engine": fnum(eng["i2_over_i1"]), "dev_pct": round(d, 4) if d is not None else None,
                                  "rms_recomputed": rc["rms_A"], "rms_engine": [fnum(eng[k]) for k in ("Ia_rms", "Ib_rms", "Ic_rms")], "harmonics": rc["harmonics_A"], "pass": d is not None and abs(d) <= TOL["recompute"] * 100})
    if recompute_tbl:
        worst = max(abs(r["dev_pct"]) for r in recompute_tbl if r["dev_pct"] is not None)
        rows_out.append(row("합성 파형 독립 재계산 I₂/I₁ (최대 편차, 전 파형)", "all", round(worst, 4), 0.0, TOL["recompute"], "%", f"{len(recompute_tbl)} 파형, 기본파 외 고조파 0", "recompute_i2i1.py", "bool", {"pass_override": worst <= TOL["recompute"] * 100}))
        rows_out[-1]["pass"] = worst <= TOL["recompute"] * 100

    # ---------------- 민감도: 고유 비대칭 각도 0/90/180 → 4턴
    for ang in ANGLES:
        run = runs[f"a{ang}"]
        if run is None:
            continue
        for sid in ("itsc_4turn", "itsc_4turn_resistive", "healthy"):
            t0, t1, _ = scene_t[sid]
            s_ = steady(run["state"], sid, t0, t1, "i2_over_i1")
            if not s_:
                continue
            final = [r for r in run["state"] if r["step"] == sid][-1]
            exp = S[sid]["i2_over_i1_composite"]["expected_by_angle"].get(f"angle_{ang}")
            sens.append({"kind": "inherent_angle", "angle_deg": ang, "scenario": sid, "i2i1_composite": round(s_["mean"], 6), "expected": exp, "dev_pct": round(pct(s_["mean"], exp), 3) if exp else None,
                         "status": final["status"], "cause": final["cause_signal"], "i2i1_level": final["lvl_i2i1"]})
    # 부하 50/100 %: I₂/I₁ vs I₂/I_n (4턴)
    for sid in ("itsc_4turn", "itsc_4turn_halfload"):
        t0, t1, step = scene_t[sid]
        s1 = steady(base["state"], sid, t0, t1, "i2_over_i1"); s2 = steady(base["state"], sid, t0, t1, "i2_over_in")
        sens.append({"kind": "load", "scenario": sid, "load": step["load"], "i2i1": round(s1["mean"], 6) if s1 else None, "i2_over_in": round(s2["mean"], 6) if s2 else None,
                     "lut_i2i1_fault": S[sid]["i2_over_i1_fault"]["value"], "lut_i2_over_in": S[sid]["i2_over_in"]["lut"]})

    # ---------------- 열: ΔT 정상 상태 (통합 모드 4·12·16턴)
    thermal_tbl = []
    for T, run in thermal.items():
        if run is None:
            continue
        sid = f"itsc_{T}turn"; R = S[sid]
        rows_ = run["state"]; sc = {s["id"]: s for s in run["meta"]["steps"]}
        t_inj = float(run["meta"]["steps"][1]["t0_s"])
        pre = [r for r in rows_ if r["step"] == "normal"]; post = [r for r in rows_ if r["step"] == sid]
        if not post:
            continue
        T0 = fnum(pre[-1]["T_winding"]) if pre else fnum(post[0]["T_winding"]); Tend = fnum(post[-1]["T_winding"]); amb = 25.0
        dT_sim = Tend - amb
        dT_ref = R["winding_dT_steady_C"]["model"]["winding"]
        ts_ = float(sc[sid].get("time_scale") or 1.0); t_sim_min = (fnum(post[-1]["t_real_s"]) - t_inj) * ts_ / 60
        thermal_tbl.append({"scenario": sid, "T_winding_before_C": T0, "T_winding_end_C": Tend, "dT_end_C": round(dT_sim, 3), "dT_steady_ref_C": dT_ref, "dev_pct": round(pct(dT_sim, dT_ref), 3),
                            "sim_minutes_after_injection": round(t_sim_min, 1), "p_loop_W": fnum(post[-1]["p_loop_total_W"]), "status_end": post[-1]["status"], "cause_end": post[-1]["cause_signal"],
                            "curve": {"t_min": [round((fnum(r["t_real_s"]) - t_inj) * ts_ / 60, 2) for r in post], "T": [fnum(r["T_winding"]) for r in post]}, "T_before_curve": [fnum(r["T_winding"]) for r in pre]})
        rows_out.append(row("권선 온도 상승 ΔT [°C] 주입 후 정상 상태 vs 열 모델 에너지 수지 역산", sid, round(dT_sim, 3), dT_ref, TOL["dT"], "°C", f"시뮬 시간 {t_sim_min:.0f} 분 후 (τ 권선 ≈ 36 분), 루프 손실 {fnum(post[-1]['p_loop_total_W'])} W", "thermal.yaml + LUT P_loop (모델 자기 일관성)"))
        ev = [e for e in run["events"] if fnum(e["t_real_s"]) >= t_inj]
        events_tbl.append({"scenario": sid + " (integrated thermal)", "t_inject_s": t_inj, "events": [{"t": fnum(e["t_real_s"]), "to": e["to"], "cause": e["cause_signal"], "value": e["cause_value"]} for e in ev], "final": post[-1]["status"], "cause": post[-1]["cause_signal"]})

    # ---------------- 저장
    n_pass = sum(1 for r in rows_out if r["pass"]); n_all = len(rows_out)
    result = {"generated": time.strftime("%Y-%m-%d %H:%M:%S"), "generator": "motor_dt_viz/tools/verify/sim_vs_analysis.py", "reference": "motor_dt_viz/data/electrical/analysis_reference.json",
              "runs": {k: (v["meta"] if v else None) for k, v in {**runs, **{f"thermal_{T}": t for T, t in thermal.items()}}.items()},
              "summary": {"pass": n_pass, "total": n_all, "fail": [r["item"] + " @ " + str(r["scenario"]) for r in rows_out if not r["pass"]]},
              "rows": rows_out, "events": events_tbl, "recompute": recompute_tbl, "sensitivity": sens, "thermal": thermal_tbl, "timeseries": timeseries, "tolerances": TOL}
    os.makedirs(OUT_DIR, exist_ok=True); os.makedirs(IMG_DIR, exist_ok=True)
    result = json.loads(json.dumps(result, ensure_ascii=False, default=_py))        # numpy 형 제거
    json.dump(result, open(os.path.join(OUT_DIR, "sim_vs_analysis.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    write_md(result); write_xlsx(result); export_csv(result, a.logs)
    if not a.no_fig:
        figures(result, ref)
    print(f"→ {OUT_DIR}/sim_vs_analysis.{{json,md,xlsx}}  PASS {n_pass}/{n_all}" + (f"  FAIL: {result['summary']['fail']}" if n_pass < n_all else ""))


def fmt(v):
    if isinstance(v, float):
        return f"{v:.5g}"
    if isinstance(v, list):
        return "~".join(fmt(x) for x in v)
    return "" if v is None else str(v)


def write_md(res):
    L = ["# 시뮬 ↔ 해석 대조표 (자동 생성)", "", f"생성 {res['generated']} · `{res['generator']}` · 기준 `{res['reference']}` · **PASS {res['summary']['pass']}/{res['summary']['total']}**", "",
         "| 항목 | 시나리오 | 시뮬(로그) | 해석 기준 | 허용 | 편차 | 판정 | 비고 | 출처 |", "|---|---|---|---|---|---|---|---|---|"]
    for r in res["rows"]:
        tol = "" if r["tolerance"] is None else (f"±{r['tolerance'] * 100:g} %" if r["kind"] == "rel" else str(r["tolerance"]))
        L.append(f"| {r['item']} | {r['scenario']} | {fmt(r['sim'])}{(' ' + r['unit']) if r['unit'] else ''} | {fmt(r['ref'])} | {tol} | {'' if r['dev_pct'] is None else f'{r['dev_pct']:+.3f} %'} | {'PASS' if r['pass'] else '**FAIL**'} | {r['note']} | {r['source']} |")
    L += ["", "## 판정 이벤트·탐지 지연", "", "| 시나리오 | 주입 [s] | 첫 WARN [s] | 첫 DANGER [s] | 지연 [s] | 장면 끝 | 원인 | i2i1 예상 |", "|---|---|---|---|---|---|---|---|"]
    for e in res["events"]:
        if "delay_s" in e:
            L.append(f"| {e['scenario']} | {e['t_inject_s']} | {fmt(e['first_WARN_s'])} | {fmt(e['first_DANGER_s'])} | {e['delay_s']} | {e['final']} | {e['cause']} | {e['expected_i2i1_level']} |")
        else:
            L.append(f"| {e['scenario']} | {e['t_inject_s']} | " + "; ".join(f"{x['t']:.1f}→{x['to']}({x['cause']})" for x in e["events"]) + f" | | | {e['final']} | {e['cause']} | |")
    L += ["", "## 독립 재계산 (합성 파형 1 s · 2 kHz, 엔진 비의존 대칭 성분)", "", "| 실행 | 파일 | 케이스 | 재계산 I₂/I₁ | 엔진 | 편차 | 고조파 h3/h1 | 판정 |", "|---|---|---|---|---|---|---|---|"]
    for r in res["recompute"]:
        L.append(f"| {r['run']} | {r['file']} | {r['case']} | {r['recomputed']:.6f} | {r['engine']:.5f} | {r['dev_pct']:+.4f} % | {r['harmonics']['h3_over_h1']} | {'PASS' if r['pass'] else 'FAIL'} |")
    L += ["", "## 민감도", "", "### 고유 비대칭 각도 0/90/180° (4턴·저항성 4턴·정상)", "", "| 각도 | 시나리오 | 합성 I₂/I₁ | 기대 | 편차 | i2i1 등급 | 종합 | 원인 |", "|---|---|---|---|---|---|---|---|"]
    for s in res["sensitivity"]:
        if s["kind"] == "inherent_angle":
            L.append(f"| {s['angle_deg']}° | {s['scenario']} | {s['i2i1_composite']:.5f} | {fmt(s['expected'])} | {fmt(s['dev_pct'])} % | {s['i2i1_level']} | {s['status']} | {s['cause']} |")
    L += ["", "### 부하 50 % / 100 % (4턴): I₂/I₁ 은 부하에 흔들리고 I₂/I_n 은 덜 흔들린다", "", "| 시나리오 | 부하 | 시뮬 I₂/I₁(합성) | 시뮬 I₂/I_n | LUT 결함분 I₂/I₁ | LUT I₂/I_n |", "|---|---|---|---|---|---|"]
    for s in res["sensitivity"]:
        if s["kind"] == "load":
            L.append(f"| {s['scenario']} | {s['load']} | {fmt(s['i2i1'])} | {fmt(s['i2_over_in'])} | {s['lut_i2i1_fault']} | {s['lut_i2_over_in']} |")
    L += ["", "## 열 (통합 모드, 주입 후 권선 온도)", "", "| 시나리오 | 주입 전 [°C] | 끝 [°C] | ΔT 시뮬 | ΔT 정상 상태 역산 | 편차 | 시뮬 분 | P_loop [W] | 끝 상태 |", "|---|---|---|---|---|---|---|---|---|"]
    for t in res["thermal"]:
        L.append(f"| {t['scenario']} | {t['T_winding_before_C']} | {t['T_winding_end_C']} | {t['dT_end_C']} | {t['dT_steady_ref_C']} | {t['dev_pct']:+.2f} % | {t['sim_minutes_after_injection']} | {t['p_loop_W']} | {t['status_end']} ({t['cause_end']}) |")
    open(os.path.join(OUT_DIR, "sim_vs_analysis.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")


def write_xlsx(res):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "대조표"
    ws.append(["항목", "시나리오", "시뮬(로그)", "해석 기준", "허용(상대)", "편차 % (수식)", "판정 (수식)", "비고", "출처", "종류"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, r in enumerate(res["rows"], start=2):
        sim = r["sim"]; refv = r["ref"]
        ws.append([r["item"], r["scenario"], sim if not isinstance(sim, list) else str(sim), refv if not isinstance(refv, list) else f"{refv[0]}~{refv[1]}", r["tolerance"], None, None, r["note"], r["source"], r["kind"]])
        if r["kind"] == "rel" and isinstance(sim, (int, float)) and isinstance(refv, (int, float)) and refv:
            ws.cell(i, 6).value = f"=(C{i}/D{i}-1)*100"
            ws.cell(i, 7).value = f'=IF(ABS(F{i})<=E{i}*100,"PASS","FAIL")'
        elif r["kind"] == "range" and isinstance(refv, list):
            ws.cell(i, 7).value = f'=IF(AND(C{i}>={refv[0]},C{i}<={refv[1]}),"PASS","FAIL")'
        elif r["kind"] == "match":
            ws.cell(i, 7).value = f'=IF(C{i}=D{i},"PASS","FAIL")'
        else:
            ws.cell(i, 7).value = "PASS" if r["pass"] else "FAIL"
        if not r["pass"]:
            for c in ws[i]:
                c.fill = PatternFill("solid", fgColor="FFE0B2")
    ws2 = wb.create_sheet("이벤트"); ws2.append(["시나리오", "주입 s", "첫 WARN s", "첫 DANGER s", "지연", "장면 끝", "원인", "i2i1 예상"])
    for e in res["events"]:
        ws2.append([e["scenario"], e["t_inject_s"], e.get("first_WARN_s"), e.get("first_DANGER_s"), json.dumps(e.get("delay_s", e.get("events")), ensure_ascii=False), e["final"], e["cause"], e.get("expected_i2i1_level")])
    ws3 = wb.create_sheet("재계산"); ws3.append(["실행", "파일", "케이스", "재계산", "엔진", "편차 % (수식)", "판정 (수식, ±0.5 %)"])
    for i, r in enumerate(res["recompute"], start=2):
        ws3.append([r["run"], r["file"], r["case"], r["recomputed"], r["engine"], f"=(D{i}/E{i}-1)*100", f'=IF(ABS(F{i})<=0.5,"PASS","FAIL")'])
    ws4 = wb.create_sheet("민감도"); ws4.append(["종류", "각도/부하", "시나리오", "합성 I₂/I₁", "기대/I₂_In", "편차 %", "등급", "종합", "원인"])
    for s in res["sensitivity"]:
        ws4.append([s["kind"], s.get("angle_deg", s.get("load")), s["scenario"], s.get("i2i1_composite", s.get("i2i1")), s.get("expected", s.get("i2_over_in")), s.get("dev_pct"), s.get("i2i1_level"), s.get("status"), s.get("cause")])
    ws5 = wb.create_sheet("열"); ws5.append(["시나리오", "주입 전", "끝", "ΔT 시뮬", "ΔT 역산", "편차 % (수식)", "판정 (±5 %)", "시뮬 분", "P_loop W"])
    for i, t in enumerate(res["thermal"], start=2):
        ws5.append([t["scenario"], t["T_winding_before_C"], t["T_winding_end_C"], t["dT_end_C"], t["dT_steady_ref_C"], f"=(D{i}/E{i}-1)*100", f'=IF(ABS(F{i})<=5,"PASS","FAIL")', t["sim_minutes_after_injection"], t["p_loop_W"]])
    wb.save(os.path.join(OUT_DIR, "sim_vs_analysis.xlsx"))


def export_csv(res, logs):
    p = os.path.join(logs, f"export_{time.strftime('%Y%m%d_%H%M%S')}.csv")
    with open(p, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["item", "scenario", "sim", "ref", "tolerance", "dev_pct", "pass", "unit", "note", "source"])
        for r in res["rows"]:
            w.writerow([r["item"], r["scenario"], fmt(r["sim"]), fmt(r["ref"]), r["tolerance"], r["dev_pct"], r["pass"], r["unit"], r["note"], r["source"]])
    res["export_csv"] = p


def figures(res, ref):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.family"] = ["Malgun Gothic", "DejaVu Sans"]; plt.rcParams["axes.unicode_minus"] = False
    warn, danger = ref["thresholds"]["i2i1_warn"], ref["thresholds"]["i2i1_danger"]
    col = {"SAFE": "#2e7d32", "WARN": "#ef6c00", "DANGER": "#c62828"}
    # 1. 시나리오별 I₂/I₁ 시계열 (한 장, 장면 구분)
    ts = res["timeseries"]
    fig, ax = plt.subplots(figsize=(12, 4.2))
    for sid, d in ts.items():
        if not d["t"]:
            continue
        ax.plot(d["t"], d["i2i1"], lw=1.2, color="#1565c0")
        for k in range(len(d["t"]) - 1):
            ax.axvspan(d["t"][k], d["t"][k + 1], color=col.get(d["status"][k], "#999"), alpha=0.12, lw=0)
        ax.axvline(d["t0"], color="#555", ls=":", lw=0.8); ax.text(d["t0"] + 0.6, 0.53, sid.replace("itsc_", ""), rotation=90, fontsize=7, va="top", ha="left", color="#333")
    ax.axhline(warn, color="#ef6c00", ls="--", lw=1, label=f"경고 {warn}"); ax.axhline(danger, color="#c62828", ls="--", lw=1, label=f"위험 {danger}")
    ax.set_xlabel("재생 시각 [s]"); ax.set_ylabel("I₂/I₁ (합성)"); ax.set_ylim(0, 0.55); ax.legend(loc="center right"); ax.set_title("전기 시나리오 I₂/I₁ 시계열 — 판정 띠(초록/주황/빨강), 점선 = 주입 시각", pad=14)
    fig.tight_layout(); fig.savefig(os.path.join(IMG_DIR, "i2i1_timeseries.png"), dpi=130); plt.close(fig)
    # 2. 시뮬 vs 해석 막대 (결함 4종 × I₂/I₁ 결함분·I_f·손실)
    scen = ["itsc_4turn", "itsc_8turn", "itsc_12turn", "itsc_16turn"]
    def get(item_prefix, sid):
        return next((r for r in res["rows"] if r["item"].startswith(item_prefix) and r["scenario"] == sid), None)
    fig, axs = plt.subplots(1, 3, figsize=(12, 3.8))
    for ax, (pref, title) in zip(axs, [("I₂/I₁ 결함분", "I₂/I₁ 결함분"), ("단락 루프 전류", "I_f [A]"), ("단락 루프 총손실", "P_loop [W]")]):
        sims, refs, labels = [], [], []
        for sid in scen:
            r = get(pref, sid)
            if r and isinstance(r["sim"], (int, float)):
                sims.append(r["sim"]); refs.append(r["ref"]); labels.append(sid.replace("itsc_", ""))
        x = range(len(labels)); ax.bar([i - 0.18 for i in x], sims, 0.36, label="시뮬(로그)", color="#1565c0"); ax.bar([i + 0.18 for i in x], refs, 0.36, label="해석(LUT/MCCM)", color="#9e9e9e")
        ax.set_xticks(list(x)); ax.set_xticklabels(labels); ax.set_title(title); ax.legend(fontsize=8)
    fig.suptitle("시뮬 vs 해석 (결함 4종)"); fig.tight_layout(); fig.savefig(os.path.join(IMG_DIR, "sim_vs_analysis_bars.png"), dpi=130); plt.close(fig)
    # 3. 권선 온도 곡선
    if res["thermal"]:
        fig, ax = plt.subplots(figsize=(8, 4))
        for t in res["thermal"]:
            ax.plot(t["curve"]["t_min"], t["curve"]["T"], label=f"{t['scenario']} (ΔT 역산 {t['dT_steady_ref_C']:.1f} °C)")
            ax.axhline(25 + t["dT_steady_ref_C"], ls=":", lw=0.8, color="#777")
        ax.set_xlabel("주입 후 시뮬레이션 시간 [분]"); ax.set_ylabel("권선 온도 [°C]"); ax.legend(fontsize=8); ax.set_title("권선 온도 (통합 모드, 점선 = 열 모델 정상 상태 역산)")
        fig.tight_layout(); fig.savefig(os.path.join(IMG_DIR, "winding_temperature.png"), dpi=130); plt.close(fig)
    # 4. 합성 파형 1 s + 스펙트럼 (4턴 post)
    wf = next((os.path.join(res["runs"]["a0"]["argv"] and "", "") for _ in []), None)
    import glob as _g
    wfs = sorted(_g.glob(os.path.join(VIZ, "logs", "verify", "electrical_a0", "waveform_*post_itsc_4turn.csv")))
    if wfs:
        import numpy as np
        rows_ = list(csv.DictReader(open(wfs[0], encoding="utf-8")))
        t_ = np.array([float(r["t_s"]) for r in rows_]); I_ = np.array([[float(r[k]) for r in rows_] for k in ("ia_A", "ib_A", "ic_A")])
        fig, (a1, a2) = plt.subplots(1, 2, figsize=(12, 3.8))
        for k, nm in enumerate("RST"):
            a1.plot(t_[:200], I_[k][:200], label=f"{nm}상")
        a1.set_xlabel("t [s]"); a1.set_ylabel("전류 [A]"); a1.set_title("합성 파형 (4턴, 2 kHz, 처음 0.1 s)"); a1.legend(fontsize=8)
        n = len(t_); fs = 1 / (t_[1] - t_[0]); F = np.fft.rfftfreq(n, 1 / fs); A = np.abs(np.fft.rfft(I_[0])) * 2 / n
        a2.semilogy(F, A + 1e-9); a2.set_xlim(0, 600); a2.set_xlabel("주파수 [Hz]"); a2.set_ylabel("|I_R| [A]"); a2.set_title("스펙트럼 — 기본파 60 Hz 단일 (RSH·측대파 없음: 해석 파형이 아니라 페이저 합성)")
        fig.tight_layout(); fig.savefig(os.path.join(IMG_DIR, "waveform_spectrum.png"), dpi=130); plt.close(fig)


if __name__ == "__main__":
    main()
